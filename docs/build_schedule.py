"""全校课程表 + 教师课表时间合并 → Excel（一门课一行）"""
import requests, json, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.stdout.reconfigure(encoding="utf-8")

# ── Login ──
resp = requests.post("https://jw.ruc.edu.cn/secService/login", json={
    "userCode": "2025202002", "password": "Wcx44773873",
    "kaptcha": "testa", "userCodeType": "ldap"
}, headers={
    "Content-Type": "application/json", "Accept": "application/json",
    "KAPTCHA-KEY-GENERATOR-REDIS": "securityKaptchaRedisServiceAdapter",
}, timeout=15)
token = resp.json()["data"]["token"]
session = resp.cookies.get("SESSION")
H = {
    "Content-Type": "application/json;charset=UTF-8",
    "Accept": "application/json, text/plain, */*",
    "token": token, "app": "PCWEB", "userrolecode": "student",
    "Cookie": f"authcode=2025202002; SESSION={session}; token=",
    "Referer": "https://jw.ruc.edu.cn/Njw2017/index.html",
}
BASE = "https://jw.ruc.edu.cn/resService/jwxtpt/v1"

# ── Step 1: 全校课程表 (3次请求) ──
print("Step 1: 课程表...")
url_s = f"{BASE}/pkgl/pkglInfo_qxxzkb/findQxxKckbKcqdList?resourceCode=XSMH0703&apiCode=jw.pkgl.pkglInfo.controller.PkqxxkbController.findQxxKckbKcqdList"
courses = []
for page in [1, 2, 3]:
    r = requests.post(url_s, json={"jczy013id": "2026-2027-1", "sctype": "zgrmdx",
        "page": {"pageIndex": page, "pageSize": 1000}}, headers=H, timeout=60)
    courses.extend(r.json()["data"]["items"])
seen = set(); unique = []
for c in courses:
    if c["id"] not in seen: seen.add(c["id"]); unique.append(c)
print(f"  {len(unique)} 门")

# ── Step 2: 教师课表 ──
print("Step 2: 教师课表...")
url_t = f"{BASE}/xsd/xsdqxxkb_info/searchTeacherkbList?resourceCode=XSMH0703&apiCode=xsdInfo.controller.XsdQxxkbController.searchTeacherkbList"
r = requests.post(url_t, json={"jczy013id": "2026-2027-1", "pkgl002id": "c13526460000WH",
    "sctype": "zgrmdx", "skls_name": "x"}, headers=H, timeout=60)

# 建索引: (课程名, 课堂名) → [时间信息, ...]
time_index = defaultdict(list)
for g in r.json()["data"]:
    if not isinstance(g, list): continue
    for item in g:
        s, e = item.get("idjkssj", 0), item.get("idjjssj", 0)
        s_str = f"{int(s)//100:02d}:{int(s)%100:02d}" if s else ""
        e_str = f"{int(e)//100:02d}:{int(e)%100:02d}" if e else ""
        days = "一二三四五六日"
        d = int(str(item.get("pksj", "0"))[0]) - 1 if item.get("pksj") else -1
        day = f"周{days[d]}" if 0 <= d < 7 else ""
        info = f"{day} {s_str}-{e_str} {item.get('js_name','')}".strip()
        info += f" [{item.get('pkzc','')}]" if item.get('pkzc') else ""
        info += f" {item.get('sjbz_name','')}" if item.get('sjbz_name','') not in ('', '无单双周') else ""
        time_index[(item.get("kc_name", ""), item.get("ktmc_name", ""))].append(info)

print(f"  索引: {len(time_index)} 个课程")

# ── Step 3: 合并写 Excel ──
print("Step 3: 写 Excel...")
wb = Workbook()
ws = wb.active
ws.title = "2026秋季课程表"

cols_def = [
    ("kcmc_name", "课程名称", 18),
    ("ktmc_name", "课堂名称", 18),
    ("skls_name", "教师", 12),
    ("kkdw_name", "开课单位", 16),
    ("zxf", "学分", 5),
    ("kclb_name", "课程类别", 10),
    ("khfs_name", "考核方式", 10),
    ("_time", "上课时间/地点", 50),
    ("jxl_name", "教学楼", 14),
    ("classroom_type", "教室类型", 10),
    ("jszc_name", "教师职称", 12),
    ("pklb_name", "排课类别", 14),
    ("jhxs", "计划学时", 8),
    ("xxrs", "限选", 5),
    ("xkrs", "已选", 5),
    ("skdx_name", "上课对象", 50),
    ("kcbh", "课程编号", 12),
    ("xnxq_name", "学期", 22),
]

hf = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
hfl = PatternFill("solid", fgColor="4472C4")
ha = Alignment(horizontal="center", vertical="center")
th = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
df = Font(name="微软雅黑", size=10)
da = Alignment(vertical="center", wrap_text=True)

for ci, (_, title, _) in enumerate(cols_def, 1):
    c = ws.cell(row=1, column=ci, value=title)
    c.font = hf; c.fill = hfl; c.alignment = ha; c.border = th

for ri, course in enumerate(unique, 2):
    cn, ktn = course.get("kcmc_name", ""), course.get("ktmc_name", "")
    times = time_index.get((cn, ktn), [])
    time_str = "\n".join(times) if times else ""

    for ci, (key, _, _) in enumerate(cols_def, 1):
        v = time_str if key == "_time" else course.get(key, "")
        if v and not isinstance(v, str): v = str(v)
        c = ws.cell(row=ri, column=ci, value=v if v else "")
        c.font = df; c.alignment = da; c.border = th

ws.auto_filter.ref = ws.dimensions
ws.freeze_panes = "A2"

for i, (_, _, w) in enumerate(cols_def):
    ws.column_dimensions[chr(65 + i)].width = w

path = r"C:\Users\wangc\Desktop\2026-2027秋季全校课程表.xlsx"
wb.save(path)

# 统计
with_time = sum(1 for c in unique if time_index.get((c.get("kcmc_name",""), c.get("ktmc_name",""))))
print(f"\nDone: {path}")
print(f"{len(unique)} 行, 其中 {with_time} 门有时间信息")
